"""Original image fixtures, bounded OCR and source-location preservation."""

import base64
import io
import json
import os

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from modulemate import ocr
from modulemate.db import UserError
from modulemate.ingest import parse_upload

TEXT = "The valve opens at 42 degrees. Do not open it earlier."


def image_bytes(format="PNG"):
    image = Image.new("RGB", (1400, 220), "white")
    ImageDraw.Draw(image).text((30, 70), TEXT, font=ImageFont.load_default(size=38), fill="black")
    output = io.BytesIO()
    image.save(output, format=format)
    return output.getvalue()


@pytest.fixture
def recognition(monkeypatch):
    monkeypatch.setattr(ocr, "local_ocr_engine", lambda: "Apple Vision")
    monkeypatch.setattr(ocr, "_vision", lambda data: ocr.Recognition(TEXT))


def mixed_pdf():
    writer = PdfWriter()
    page = writer.add_page(PdfReader(io.BytesIO(image_bytes("PDF"))).pages[0])
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page["/Resources"][NameObject("/Font")] = DictionaryObject(
        {NameObject("/F1"): writer._add_object(font)}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        page.get_contents().get_data()
        + b"\nBT /F1 14 Tf 10 200 Td (Native maintenance instruction.) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def vector_pdf():
    # Original block-letter outlines: no PDF text operators or image objects.
    glyphs = {
        "O": ("01110", "10001", "10001", "10001", "10001", "10001", "01110"),
        "P": ("11110", "10001", "10001", "11110", "10000", "10000", "10000"),
        "E": ("11111", "10000", "10000", "11110", "10000", "10000", "11111"),
        "N": ("10001", "11001", "11001", "10101", "10011", "10011", "10001"),
        "A": ("01110", "10001", "10001", "11111", "10001", "10001", "10001"),
        "T": ("11111", "00100", "00100", "00100", "00100", "00100", "00100"),
        "4": ("00010", "00110", "01010", "10010", "11111", "00010", "00010"),
        "2": ("01110", "10001", "00001", "00010", "00100", "01000", "11111"),
        " ": ("00000",) * 7,
    }
    paths = [
        f"{20 + (i * 6 + x) * 6} {110 - y * 6} 6 6 re f"
        for i, char in enumerate("OPEN AT 42")
        for y, row in enumerate(glyphs[char])
        for x, bit in enumerate(row)
        if bit == "1"
    ]
    writer = PdfWriter()
    page = writer.add_blank_page(width=450, height=180)
    stream = DecodedStreamObject()
    stream.set_data("\n".join(paths).encode())
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_pdf_without_native_text_or_image_objects_still_gets_ocr(recognition):
    data = vector_pdf()
    page = PdfReader(io.BytesIO(data)).pages[0]
    assert not page.extract_text() and not len(page.images)
    parsed = parse_upload("Outlined label.pdf", data, isolated=False)
    assert TEXT in parsed.text and parsed.chunks[0].locator == "page 1 · OCR"


def test_mixed_pdf_keeps_native_text_and_recovers_image_text(recognition):
    parsed = parse_upload("Valve.pdf", mixed_pdf(), isolated=False)
    assert "Native maintenance instruction." in parsed.text and TEXT in parsed.text
    assert {p.locator for p in parsed.chunks} == {"page 1", "page 1 · OCR"}


def test_reimporting_the_same_original_can_add_previously_missing_ocr(recognition, store, course):
    data = mixed_pdf()
    old = parse_upload("Valve.pdf", data, isolated=False, use_ocr=False)
    updated = parse_upload("Valve.pdf", data, isolated=False)
    first, created = store.add_document(course, old)
    second, added = store.add_document(course, updated)
    assert created and added and first != second
    assert old.digest != updated.digest and TEXT not in store.document(first)["text"]
    assert TEXT in store.document(second)["text"]
    assert store.add_document(course, updated) == (second, False)


def test_nested_slide_image_and_notes_keep_slide_locator(recognition):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    group.shapes.add_picture(io.BytesIO(image_bytes()), Inches(1), Inches(1), width=Inches(8))
    slide.notes_slide.notes_text_frame.text = "Read the valve label first."
    buffer = io.BytesIO()
    presentation.save(buffer)
    parsed = parse_upload("Valve.pptx", buffer.getvalue(), isolated=False)
    assert "Read the valve label first." in parsed.text and TEXT in parsed.text
    assert any(p.locator == "slide 1 · image 1 · OCR" for p in parsed.chunks)


def test_docx_table_picture_is_not_omitted(recognition):
    document = Document()
    document.add_paragraph("Native instructions.")
    cell = document.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes()))
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_upload("Valve.docx", buffer.getvalue(), isolated=False)
    assert "Native instructions." in parsed.text and TEXT in parsed.text
    assert any("OCR" in p.locator for p in parsed.chunks)


def test_docx_header_picture_is_read_with_its_header_location(recognition):
    document = Document()
    document.add_paragraph("Native instructions.")
    document.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(image_bytes()))
    buffer = io.BytesIO()
    document.save(buffer)
    parsed = parse_upload("Valve.docx", buffer.getvalue(), isolated=False)
    assert TEXT in parsed.text
    assert any(p.locator.startswith("header default") and "OCR" in p.locator for p in parsed.chunks)


def test_notebook_saved_text_images_and_attachments_are_static_evidence(recognition, tmp_path):
    target = tmp_path / "must-not-exist"
    image = base64.b64encode(image_bytes()).decode()
    data = {
        "cells": [
            {
                "cell_type": "markdown",
                "source": "Original diagram.",
                "attachments": {"valve.png": {"image/png": image}},
            },
            {
                "cell_type": "code",
                "source": f"open({str(target)!r},'w').write('executed')",
                "outputs": [
                    {
                        "data": {
                            "text/plain": ["Saved valve reading: 42 degrees."],
                            "image/png": image,
                        }
                    },
                ],
            },
        ]
    }
    parsed = parse_upload("Valve.ipynb", json.dumps(data).encode(), isolated=False)
    assert not target.exists()
    assert TEXT in parsed.text and "Saved valve reading: 42 degrees." in parsed.text
    assert {
        "cell 1 · attachment 1 · OCR",
        "cell 2 · saved output 1",
        "cell 2 · saved output 1 · OCR",
    } <= {p.locator for p in parsed.chunks}
    assert any("nothing is executed" in warning for warning in parsed.warnings)


def test_notebook_html_and_malformed_images_do_not_execute_or_hide_omission(recognition):
    data = {
        "cells": [
            {
                "cell_type": "code",
                "source": "# Original static example",
                "outputs": [
                    {"data": {"text/html": "<script>fetch('https://example.invalid')</script>"}},
                    {"data": {"image/png": "invalid base64!"}},
                ],
            }
        ]
    }
    parsed = parse_upload("Static.ipynb", json.dumps(data).encode(), isolated=False)
    assert "fetch" not in parsed.text
    assert any("rich output is omitted" in warning for warning in parsed.warnings)
    assert any("cell 1 · saved output 2" in warning for warning in parsed.warnings)


def test_ocr_failure_preserves_native_pdf_and_discloses_missing_image_text(monkeypatch):
    monkeypatch.setattr(ocr, "local_ocr_engine", lambda: None)
    parsed = parse_upload("Valve.pdf", mixed_pdf(), isolated=False)
    assert "Native maintenance instruction." in parsed.text and TEXT not in parsed.text
    assert any("Unread image locations" in warning for warning in parsed.warnings)


def test_failed_image_ocr_is_actionable(monkeypatch):
    monkeypatch.setattr(ocr, "local_ocr_engine", lambda: None)
    with pytest.raises(UserError, match="OCR is unavailable"):
        parse_upload("Valve.png", image_bytes(), isolated=False)


def test_ocr_bounds_and_low_confidence_are_visible_without_discarding_text(
    recognition, monkeypatch
):
    monkeypatch.setattr(ocr, "_vision", lambda data: ocr.Recognition(TEXT, uncertain=True))
    session = ocr.OCRSession()
    assert session.read(image_bytes(), "slide 1") == TEXT
    assert any("low recognition confidence" in w for w in session.warnings())
    session.attempts = ocr.MAX_IMAGES
    assert (
        session.read(image_bytes(), "slide 2") == TEXT
    )  # Reuse recognized text within this import.
    assert session.read(image_bytes() + b"\n", "slide 3") == ""
    assert any("slide 3" in w and "resource limit" in w for w in session.warnings())


def test_normal_65_slide_deck_does_not_stop_after_32_images(recognition):
    session = ocr.OCRSession()
    for i in range(65):
        assert session.read(image_bytes() + bytes([i]), f"slide {i + 1}") == TEXT
    assert session.attempts == 65 and not session.skipped


def test_native_duplicates_are_removed_without_fuzzy_number_matching():
    assert ocr.additional_text("Valve label\n42 degrees", "Valve label") == "42 degrees"
    assert ocr.additional_text("42 degrees", "142 degrees") == "42 degrees"
    assert ocr.additional_text("pH = 7\npH = 7", "") == "pH = 7"


@pytest.mark.skipif(os.environ.get("MODULEMATE_TEST_OCR") != "1", reason="Opt-in real local OCR")
@pytest.mark.parametrize("kind", ["image", "scan", "mixed", "vector"])
def test_real_offline_ocr_in_isolated_parser(kind):
    assert ocr.local_ocr_engine(), "Install the documented local OCR engine"
    name, data = {
        "image": ("Valve.png", image_bytes()),
        "scan": ("Valve.pdf", image_bytes("PDF")),
        "mixed": ("Valve.pdf", mixed_pdf()),
        "vector": ("Outlined label.pdf", vector_pdf()),
    }[kind]
    parsed = parse_upload(name, data)
    assert ("OPEN AT 42" if kind == "vector" else TEXT) in parsed.text
    assert any("OCR" in p.locator for p in parsed.chunks)
    if kind == "mixed":
        assert "Native maintenance instruction." in parsed.text


@pytest.mark.skipif(
    os.environ.get("MODULEMATE_TEST_OCR_STRESS") != "1", reason="Opt-in 65-page local OCR"
)
def test_real_65_page_scanned_document_keeps_its_final_page():
    images = []
    for i in range(65):
        image = Image.new("RGB", (1000, 250), "white")
        ImageDraw.Draw(image).text(
            (30, 40),
            f"Original scan page {i + 1}.\nThe valve opens at 42 degrees.",
            font=ImageFont.load_default(size=32),
            fill="black",
        )
        images.append(image)
    buffer = io.BytesIO()
    images[0].save(buffer, format="PDF", save_all=True, append_images=images[1:])
    parsed = parse_upload("Original scans.pdf", buffer.getvalue())
    assert len({p.locator for p in parsed.chunks}) == 65
    assert "Original scan page 65." in parsed.text
