"""Actual original Office packages for previously omitted static evidence."""

import io

from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.package import Part as SlidePart
from pptx.opc.packuri import PackURI as SlideURI
from pptx.oxml import parse_xml as slide_xml
from pptx.util import Inches

from modulemate.ingest import parse_upload


def save(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_word_referenced_headers_footer_and_text_box_are_searchable():
    document = Document()
    document.add_paragraph("Native body instruction.")
    document.sections[0].header.paragraphs[0].text = "The inspection edition is Cedar."
    document.sections[0].footer.paragraphs[0].text = "Do not reuse a damaged seal."
    picture = parse_xml("""<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:v="urn:schemas-microsoft-com:vml"><v:shape style="width:180pt;height:40pt">
        <v:textbox><w:txbxContent><w:p><w:r><w:t>Boxed limit: six tiles.</w:t></w:r></w:p>
        </w:txbxContent></v:textbox></v:shape></w:pict>""")
    document.paragraphs[0].add_run()._r.append(picture)
    parsed = parse_upload("Instruction.docx", save(document))
    assert all(t in parsed.text for t in ("Cedar", "damaged seal", "six tiles", "Native body"))
    assert any(p.locator.startswith("header default") for p in parsed.chunks)
    assert any(p.locator == "text box 1" for p in parsed.chunks)


def test_only_referenced_word_notes_are_read():
    document = Document()
    paragraph = document.add_paragraph("Check the attached note.")
    root = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    for kind in ("footnote", "endnote"):
        part = Part(
            PackURI(f"/word/{kind}s.xml"),
            f"application/vnd.openxmlformats-officedocument.wordprocessingml.{kind}s+xml",
            f'''<w:{kind}s xmlns:w="{root}"><w:{kind} w:id="7"><w:p><w:r><w:t>
            {kind}: keep two spare seals.</w:t></w:r></w:p></w:{kind}>
            <w:{kind} w:id="99"><w:p><w:r><w:t>ORPHAN NOTE</w:t></w:r></w:p></w:{kind}>
            </w:{kind}s>'''.encode(),
            document.part.package,
        )
        document.part.relate_to(
            part, f"http://schemas.openxmlformats.org/officeDocument/2006/relationships/{kind}s"
        )
        reference = OxmlElement(f"w:{kind}Reference")
        reference.set(qn("w:id"), "7")
        paragraph.add_run()._r.append(reference)
    parsed = parse_upload("Notes.docx", save(document))
    assert {"footnote 7", "endnote 7"} <= {p.locator for p in parsed.chunks}
    assert "keep two spare seals" in parsed.text and "ORPHAN NOTE" not in parsed.text


def test_native_chart_pairs_categories_with_raw_values_and_preserves_format():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = CategoryChartData(number_format="0%")
    data.categories = ["Front", "Back"]
    data.add_series("Rough tiles", [0.25, 0.6])
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(7), Inches(4), data
    )
    parsed = parse_upload("Inspection.pptx", save(presentation))
    assert "Point 0: Category = Front; Value = 0.25" in parsed.text
    assert "Point 1: Category = Back; Value = 0.6" in parsed.text
    assert "Value number format: 0%" in parsed.text
    assert "Rough tiles" in parsed.text
    assert any("slide 1 · chart 1" in p.locator for p in parsed.chunks)


def test_native_scatter_chart_does_not_mislabel_x_values_as_categories():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    data = XyChartData()
    series = data.add_series("Trial")
    series.add_data_point(2, 7)
    series.add_data_point(5, 11)
    slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER, Inches(1), Inches(1), Inches(7), Inches(4), data
    )
    parsed = parse_upload("Trial.pptx", save(presentation))
    assert "Point 0: X = 2; Y = 7" in parsed.text
    assert "Point 1: X = 5; Y = 11" in parsed.text


def test_smartart_reads_referenced_labels_without_inventing_relationships():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    xml = b"""<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>
        <dgm:pt modelId="1"><dgm:t><a:p><a:r><a:t>Collect</a:t></a:r></a:p></dgm:t></dgm:pt>
        <dgm:pt modelId="2"><dgm:t><a:p><a:r><a:t>Review</a:t></a:r></a:p></dgm:t></dgm:pt>
        </dgm:ptLst></dgm:dataModel>"""
    part = SlidePart(
        SlideURI("/ppt/diagrams/data1.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml",
        presentation.part.package,
        xml,
    )
    relation_type = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/diagramData"
    )
    rid = slide.part.relate_to(part, relation_type)
    frame = slide_xml(f'''<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
        <p:nvGraphicFramePr><p:cNvPr id="2" name="Process labels"/><p:cNvGraphicFramePr/><p:nvPr/></p:nvGraphicFramePr>
        <p:xfrm><a:off x="0" y="0"/><a:ext cx="1000000" cy="1000000"/></p:xfrm>
        <a:graphic><a:graphicData uri="http://schemas.openxmlformats.org/drawingml/2006/diagram">
        <dgm:relIds r:dm="{rid}"/></a:graphicData></a:graphic></p:graphicFrame>''')
    slide.shapes._spTree.append(frame)
    parsed = parse_upload("Process.pptx", save(presentation))
    assert "Collect" in parsed.text and "Review" in parsed.text
    assert any("SmartArt 1 labels" in p.locator for p in parsed.chunks)
    assert "arrow" not in parsed.text
