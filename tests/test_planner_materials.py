"""Date binding through actual document extraction, not only hand-built strings."""

import io
import json
import os

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from modulemate.ingest import parse_upload
from modulemate.planner import candidates
from modulemate.search import Passage

TEXT = "Final exam: November 2, 2026, 14:00-15:30 KST."


def payload(kind):
    output = io.BytesIO()
    if kind == "docx":
        doc = Document()
        doc.add_heading("Assessment dates", 1)
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = TEXT
        doc.save(output)
    elif kind == "pptx":
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = "Assessment dates"
        slide.placeholders[1].text = TEXT
        deck.save(output)
    elif kind == "pdf":
        pdf = PdfWriter()
        page = pdf.add_blank_page(width=600, height=200)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})}
        )
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 12 Tf 20 150 Td (" + TEXT + ") Tj ET").encode())
        page[NameObject("/Contents")] = pdf._add_object(stream)
        pdf.write(output)
    elif kind == "ipynb":
        return json.dumps({"cells": [{"cell_type": "markdown", "source": TEXT}]}).encode()
    else:
        return TEXT.encode()
    return output.getvalue()


@pytest.mark.parametrize("kind", ["txt", "md", "docx", "pptx", "pdf", "ipynb"])
def test_actual_native_import_preserves_date_time_and_source_locator(kind):
    parsed = parse_upload("Assessment." + kind, payload(kind), isolated=False, use_ocr=False)
    passages = [
        Passage(i, "one", parsed.name, c.locator, c.text) for i, c in enumerate(parsed.chunks, 1)
    ]
    blocks = candidates(passages)
    assert any(d["date"] == "2026-11-02" for b in blocks for d in b["dates"])
    assert any([t["time"] for t in b["times"]] == ["14:00", "15:30"] for b in blocks)
    assert all(b["sources"][0]["locator"] for b in blocks)
    assert all(
        " ".join(b["sources"][0]["quote"].split())
        in " ".join(passages[b["sources"][0]["passage_id"] - 1].text.split())
        for b in blocks
    )


@pytest.mark.skipif(
    os.environ.get("MODULEMATE_TEST_OCR") != "1", reason="Opt-in real local date OCR"
)
@pytest.mark.parametrize("kind", ["png", "jpeg", "pdf"])
def test_actual_ocr_schedule_keeps_calendar_numbers(kind):
    image = Image.new("RGB", (1900, 300), "white")
    ImageDraw.Draw(image).text((30, 90), TEXT, font=ImageFont.load_default(size=44), fill="black")
    output = io.BytesIO()
    image.save(output, format=kind.upper())
    parsed = parse_upload("Scanned schedule." + kind, output.getvalue())
    passages = [
        Passage(i, "one", parsed.name, c.locator, c.text) for i, c in enumerate(parsed.chunks, 1)
    ]
    blocks = candidates(passages)
    assert any(d["date"] == "2026-11-02" for b in blocks for d in b["dates"]), parsed.text
    assert any([t["time"] for t in b["times"]] == ["14:00", "15:30"] for b in blocks), parsed.text
    assert any("OCR" in p.locator for p in passages)
