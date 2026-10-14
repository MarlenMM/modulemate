import io
import json
import zipfile

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from modulemate.db import UserError
from modulemate.ingest import MAX_BYTES, parse_upload, safe_name, split_text


def test_text_short_unicode_and_paths():
    parsed = parse_upload("../../notes.txt", "Quiz: 금요일.\n概率是可能性的度量。".encode())
    assert parsed.name == "notes.txt"
    assert "概率" in parsed.chunks[0].text
    assert parsed.chunks[0].locator == "part 1"
    assert safe_name("..\\evil\x00.md") == "evil.md"


def test_chunking_no_gaps_and_bounded():
    text = "".join(chr(0x4E00 + i) for i in range(5000))
    pieces = split_text(text)
    assert all(len(p) <= 1200 for p in pieces)
    assert set("".join(pieces)) == set(text)
    assert pieces[0][-150:] == pieces[1][:150]
    with pytest.raises(ValueError):
        split_text(text, size=100, overlap=100)


@pytest.mark.parametrize(
    "name,data",
    [
        ("bad.exe", b"content"),
        ("bad.txt", b""),
        ("bad.txt", b"\x00hello"),
        ("bad.txt", b"\xff"),
        ("bad.pdf", b"not pdf"),
        ("bad.docx", b"not zip"),
        ("bad.ipynb", b"{}"),
        ("bad.ipynb", b"null"),
    ],
)
def test_rejects_malformed_uploads(name, data):
    with pytest.raises(UserError):
        parse_upload(name, data)


def test_size_limit():
    with pytest.raises(UserError, match="15 MB"):
        parse_upload("large.txt", b"a" * (MAX_BYTES + 1))


def test_notebook_ignores_outputs_and_never_executes(tmp_path):
    target = tmp_path / "should-not-exist"
    data = {
        "cells": [
            {"cell_type": "markdown", "source": ["Original explanation."]},
            {
                "cell_type": "code",
                "source": f"open({str(target)!r},'w').write('oops')",
                "outputs": [{"text": "SECRET OUTPUT"}],
            },
        ]
    }
    parsed = parse_upload("demo.ipynb", json.dumps(data).encode())
    assert not target.exists()
    assert "SECRET OUTPUT" not in parsed.text
    assert {c.locator for c in parsed.chunks} == {"cell 1", "cell 2"}


def test_docx_preserves_paragraph_table_order():
    doc = Document()
    doc.add_paragraph("Before the table.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Inside the table."
    doc.add_paragraph("After the table.")
    data = io.BytesIO()
    doc.save(data)
    parsed = parse_upload("sample.docx", data.getvalue())
    assert parsed.text.index("Before") < parsed.text.index("Inside") < parsed.text.index("After")
    assert parsed.warnings


def test_pptx_extracts_slide_and_notes():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Original biology example"
    slide.placeholders[1].text = "Mitochondria participate in cellular respiration."
    slide.notes_slide.notes_text_frame.text = "Remember the source location."
    data = io.BytesIO()
    prs.save(data)
    parsed = parse_upload("sample.pptx", data.getvalue())
    assert parsed.chunks[0].locator == "slide 1"
    assert "Remember the source location" in parsed.text


def test_pdf_page_locations_and_empty_page_warning():
    writer = PdfWriter()
    page = writer.add_blank_page(width=400, height=400)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 10 300 Td (Original PDF text about photosynthesis.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    writer.add_blank_page(width=400, height=400)
    data = io.BytesIO()
    writer.write(data)
    parsed = parse_upload("sample.pdf", data.getvalue())
    assert parsed.chunks[0].locator == "page 1"
    assert "photosynthesis" in parsed.text
    assert "1 pages have no readable text" in parsed.warnings[0]


def test_encrypted_pdf_rejected():
    writer = PdfWriter()
    writer.add_blank_page(width=400, height=400)
    writer.encrypt("password")
    data = io.BytesIO()
    writer.write(data)
    with pytest.raises(UserError, match="Unlock"):
        parse_upload("secret.pdf", data.getvalue())


def test_office_entity_and_expansion_rejected():
    for content in [
        b'<!DOCTYPE doc [<!ENTITY x SYSTEM "file:///etc/passwd">]>',
        b"a" * (13 * 1024 * 1024),
    ]:
        data = io.BytesIO()
        with zipfile.ZipFile(data, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", content)
        with pytest.raises(UserError):
            parse_upload("unsafe.docx", data.getvalue())


def test_prose_overlap_does_not_start_mid_word():
    text = " ".join("photosynthesis mitochondria respiration" for _ in range(120))
    pieces = split_text(text)
    assert all(p.split()[0] in {"photosynthesis", "mitochondria", "respiration"} for p in pieces)
    assert all(len(p) <= 1200 for p in pieces)
