"""Bounded extraction of user uploads, never imported/executed as code."""

from __future__ import annotations

import base64
import binascii
import hashlib
import io
import json
import os
import re
import subprocess
import sys
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

from modulemate.db import UserError
from modulemate.document_parts import (
    chart_text,
    embedded_images,
    smartart_text,
    word_additional_parts,
)
from modulemate.ocr import MAX_SIDE, OCRSession

MAX_BYTES = 15 * 1024 * 1024
MAX_CHARACTERS = 2_000_000
MAX_UNITS = 500
MAX_ARCHIVE_BYTES = 60 * 1024 * 1024
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".webp"}
SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".pptx", ".ipynb"} | IMAGE_TYPES


@dataclass(frozen=True)
class Chunk:
    locator: str
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    name: str
    digest: str
    text: str
    chunks: list[Chunk]
    warnings: list[str]


def safe_name(name: str) -> str:
    name = unicodedata.normalize("NFKC", name.replace("\\", "/").split("/")[-1])
    name = "".join(c for c in name if not unicodedata.category(c).startswith("C"))
    return name.strip()[:180] or "Untitled.txt"


def split_text(text: str, size: int = 1200, overlap: int = 150) -> list[str]:
    if size <= 0 or not 0 <= overlap < size:
        raise ValueError("Chunk overlap must be smaller than the positive chunk size.")
    text = text.replace("\x00", "").replace("\r\n", "\n").strip()
    result = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            boundary = max(
                text.rfind("\n", start + size // 2, end), text.rfind(" ", start + size // 2, end)
            )
            if boundary > start:
                end = boundary
        piece = text[start:end].strip()
        if piece:
            result.append(piece)
        if end == len(text):
            break
        candidate = max(start + 1, end - overlap)
        if candidate > 0 and text[candidate - 1].isalnum() and text[candidate].isalnum():
            boundary = max(
                text.rfind(" ", max(start + 1, candidate - 50), candidate + 1),
                text.rfind("\n", max(start + 1, candidate - 50), candidate + 1),
            )
            if boundary > start:
                candidate = boundary + 1
        start = candidate
    return result


def structured_chunks(text: str) -> list[tuple[str, str]]:
    """Keep Markdown sections together and carry their real heading hierarchy.

    Headings inside fenced code are code. A small section is never glued to the
    next unrelated concept merely to fill a character window.
    """
    sections, body, headings = [], [], []
    fence = None

    def flush():
        content = "\n".join(body).strip()
        if not content:
            return
        title = " › ".join(value for _, value in headings)
        prefix = "\n".join("#" * level + " " + value for level, value in headings)
        for piece in split_text(content):
            sections.append((title, (prefix + "\n" + piece).strip()))

    for line in text.splitlines():
        marker = re.match(r"^\s*(`{3,}|~{3,})", line)
        if marker:
            kind = marker[1][0]
            fence = None if fence == kind else (kind if fence is None else fence)
        match = re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line) if fence is None else None
        if match:
            flush()
            body = []
            level = len(match[1])
            headings = [(n, title) for n, title in headings if n < level]
            headings.append((level, match[2]))
        else:
            body.append(line)
    flush()
    # Preserve a heading-only note instead of silently producing an empty file.
    return sections or [("", part) for part in split_text(text)]


def _check_archive(data: bytes, extension: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 2500 or sum(i.file_size for i in entries) > MAX_ARCHIVE_BYTES:
                raise UserError("Office file expands beyond the safe processing limit.")
            required = "word/document.xml" if extension == ".docx" else "ppt/presentation.xml"
            if required not in archive.namelist():
                raise UserError("The file does not match its Office document extension.")
            for item in entries:
                if item.flag_bits & 1 or item.file_size > 12 * 1024 * 1024:
                    raise UserError(
                        "Encrypted or unusually large Office components are not supported."
                    )
                if item.filename.lower().endswith((".xml", ".rels")):
                    xml = archive.read(item).upper()
                    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
                        raise UserError("Office XML containing external entities is not supported.")
    except (zipfile.BadZipFile, RuntimeError, KeyError) as exc:
        raise UserError("This Office file is damaged or unsupported.") from exc


def _extract(
    data: bytes, extension: str, *, use_ocr: bool = True
) -> tuple[list[tuple[str, str]], list[str]]:
    """Runs in a resource-limited subprocess in normal application use."""
    warnings = []
    units = []
    ocr = OCRSession(enabled=use_ocr)
    if extension in {".txt", ".md"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise UserError("Save text files as UTF-8, then upload again.") from exc
        if "\x00" in text:
            raise UserError("This looks like a binary file. Upload UTF-8 text instead.")
        # Splitting before recognizing sections used to orphan a heading from
        # the definition it governs. Chunk only after parsing the structure.
        units = [("part 1", text)]
    elif extension == ".ipynb":
        try:
            notebook = json.loads(data)
            cells = notebook["cells"]
            if not isinstance(cells, list) or len(cells) > MAX_UNITS:
                raise UserError("Notebooks may contain at most 500 cells.")

            def notebook_text(value):
                value = "".join(value) if isinstance(value, list) else value
                if not isinstance(value, str):
                    raise UserError("Invalid notebook text.")
                return value

            def notebook_images(bundle, location):
                if any(mime.startswith("image/") for mime in bundle) and not any(
                    mime in bundle for mime in ("image/png", "image/jpeg", "image/webp")
                ):
                    ocr.skipped.append(location + " · unsupported image")
                for mime in ("image/png", "image/jpeg", "image/webp"):
                    if mime not in bundle:
                        continue
                    image_location = location + " · OCR"
                    try:
                        encoded = notebook_text(bundle[mime])
                        image = base64.b64decode(re.sub(r"\s", "", encoded), validate=True)
                        if len(image) > MAX_BYTES:
                            raise ValueError("Image limit")
                        text = ocr.read(image, image_location)
                        if text:
                            units.append((image_location, text))
                    except (ValueError, binascii.Error):
                        ocr.skipped.append(image_location)
                    break  # Alternate encodings describe the same displayed image.

            output_count = 0
            for i, cell in enumerate(cells, 1):
                if cell.get("cell_type") not in {"markdown", "code"}:
                    continue
                source = notebook_text(cell.get("source", ""))
                if source.strip():
                    units.append((f"cell {i}", source))
                attachments = cell.get("attachments", {})
                if not isinstance(attachments, dict):
                    raise UserError("Invalid notebook attachments.")
                for j, bundle in enumerate(attachments.values(), 1):
                    if isinstance(bundle, dict):
                        notebook_images(bundle, f"cell {i} · attachment {j}")
                outputs = cell.get("outputs", [])
                if not isinstance(outputs, list):
                    raise UserError("Invalid notebook outputs.")
                output_count += len(outputs)
                if len(outputs) > 100 or output_count > 2000:
                    raise UserError(
                        "Notebooks may contain at most 100 saved outputs per cell and 2,000 in total. Split this notebook."
                    )
                for j, output in enumerate(outputs, 1):
                    if not isinstance(output, dict):
                        raise UserError("Invalid notebook output.")
                    bundle = output.get("data", {})
                    if not isinstance(bundle, dict):
                        raise UserError("Invalid notebook output data.")
                    # Saved text is data, never an instruction to execute the
                    # corresponding code or recalculate the stored result.
                    text = notebook_text(bundle.get("text/plain", output.get("text", "")))
                    location = f"cell {i} · saved output {j}"
                    if text.strip():
                        units.append((location, text))
                    notebook_images(bundle, location)
                    if (
                        not text.strip()
                        and bundle
                        and not any(k.startswith("image/") for k in bundle)
                    ):
                        warnings.append(
                            f"{location} has no plain text or supported raster image; rich output is omitted."
                        )
            warnings.append(
                "Notebook code and saved outputs are read as static material; nothing is executed or recomputed. Saved outputs may be stale. Images/attachments use local OCR; interactive widgets and HTML-only output need supplied text."
            )
        except (ValueError, KeyError, TypeError, AttributeError) as exc:
            if isinstance(exc, UserError):
                raise
            raise UserError("This is not a valid notebook.") from exc
    elif extension == ".pdf":
        from pypdf import PdfReader

        if not data.lstrip().startswith(b"%PDF-"):
            raise UserError("The file does not contain a PDF header.")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise UserError("Unlock the PDF before uploading it.")
        if len(reader.pages) > MAX_UNITS:
            raise UserError("PDFs may contain at most 500 pages. Split this file first.")
        empty, rendered = [], None
        try:
            for i, page in enumerate(reader.pages, 1):
                text = page.extract_text() or ""
                if text.strip():
                    units.append((f"page {i}", text))
                # A page without native text can contain outlined/vector letters
                # and no raster image objects. Mixed text/image pages need OCR
                # too. Inspect identities without decoding every embedded bitmap.
                recognized = ""
                if not text.strip() or len(page.images):
                    location = f"page {i} · OCR"
                    if not ocr.available:
                        ocr.skipped.append(location)
                    else:
                        try:
                            import pypdfium2 as pdfium

                            if rendered is None:
                                rendered = pdfium.PdfDocument(data)
                            raster = rendered[i - 1]
                            try:
                                width, height = raster.get_size()
                                scale = min(2.5, MAX_SIDE / max(width, height))
                                bitmap = raster.render(scale=scale)
                                try:
                                    image = bitmap.to_pil()
                                    buffer = io.BytesIO()
                                    image.save(buffer, format="PNG")
                                finally:
                                    bitmap.close()
                            finally:
                                raster.close()
                            recognized = ocr.read(buffer.getvalue(), location, text)
                            if recognized:
                                units.append((location, recognized))
                        except (ValueError, RuntimeError, OSError):
                            ocr.skipped.append(location)
                if not text.strip() and not recognized:
                    empty.append(i)
        finally:
            if rendered is not None:
                rendered.close()
        if empty:
            warnings.append(
                f"{len(empty)} pages have no readable text (first pages: {', '.join(map(str, empty[:10]))}). Blank pages, diagrams without text and unreadable scans need a text description."
            )
    elif extension in {".docx", ".pptx"}:
        _check_archive(data, extension)
        if extension == ".docx":
            from docx import Document
            from docx.table import Table

            doc = Document(io.BytesIO(data))
            lines = []
            # Preserve paragraph/table order rather than moving tables to the end.
            image_number = 0
            for block in doc.iter_inner_content():
                if isinstance(block, Table):
                    lines.append(
                        "\n".join(" | ".join(cell.text for cell in row.cells) for row in block.rows)
                    )
                else:
                    heading = re.match(r"Heading ([1-6])", block.style.name or "")
                    lines.append(("#" * int(heading[1]) + " " if heading else "") + block.text)
                # Images may occur inside paragraphs or table cells. Preserve
                # the surrounding block's order and inspect only embedded parts;
                # external relationships are never fetched.
                for image in embedded_images(block._element, doc.part):
                    image_number += 1
                    image_location = f"image {image_number} · OCR"
                    image_text = ocr.read(image, image_location, "\n".join(lines))
                    if image_text:
                        units.append((image_location, image_text))
            units = [
                (f"section {i}", text)
                for i, text in enumerate(split_text("\n".join(lines), 5000, 0), 1)
                if text.strip()
            ] + units
            for location, text, root, part in word_additional_parts(doc):
                if text:
                    units.append((location, text))
                if location.startswith("text box "):
                    continue  # Its embedded pictures were already read with the body.
                for j, image in enumerate(embedded_images(root, part), 1):
                    image_location = f"{location} · image {j} · OCR"
                    image_text = ocr.read(image, image_location, text)
                    if image_text:
                        units.append((image_location, image_text))
            warnings.append(
                "DOCX locations are sections, referenced headers/footers, note identifiers and text boxes, not page numbers. Conditional header/footer variants are labeled; visual layout is not reconstructed."
            )
        else:
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(data))
            if len(presentation.slides) > MAX_UNITS:
                raise UserError("Presentations may contain at most 500 slides.")

            def shape_text(shape, image_parts, charts):
                if hasattr(shape, "shapes"):
                    return "\n".join(
                        shape_text(child, image_parts, charts) for child in shape.shapes
                    )
                if shape.has_chart:
                    charts.append(chart_text(shape.chart))
                if hasattr(shape, "image"):
                    image_parts.append(shape.image.blob)
                if shape.has_table:
                    return "\n".join(
                        " | ".join(cell.text for cell in row.cells) for row in shape.table.rows
                    )
                return shape.text if shape.has_text_frame else ""

            missing = 0
            for i, slide in enumerate(presentation.slides, 1):
                images, charts = [], []
                text = "\n".join(shape_text(shape, images, charts) for shape in slide.shapes)
                if slide.has_notes_slide:
                    text += "\n" + slide.notes_slide.notes_text_frame.text
                if text.strip():
                    units.append((f"slide {i}", text))
                for j, text_chart in enumerate(charts, 1):
                    units.append((f"slide {i} · chart {j}", text_chart))
                diagram_labels = list(smartart_text(slide))
                for j, labels in enumerate(diagram_labels, 1):
                    units.append((f"slide {i} · SmartArt {j} labels", labels))
                recognized = bool(charts or diagram_labels)
                for j, image in enumerate(images, 1):
                    location = f"slide {i} · image {j} · OCR"
                    image_text = ocr.read(image, location, text)
                    if image_text:
                        units.append((location, image_text))
                        recognized = True
                if not text.strip() and not recognized:
                    missing += 1
            warnings.append(
                "Embedded image text uses local OCR. Native charts expose cached raw values and stored formats, without recalculating a workbook. SmartArt exposes text labels only; diagram relationships, visual chart readings and handwriting may need a description."
            )
            if missing:
                warnings.append(f"{missing} slides have no readable text.")
    elif extension in IMAGE_TYPES:
        text = ocr.read(data, "image 1 · OCR")
        if text:
            units.append(("image 1 · OCR", text))
    warnings.extend(ocr.warnings())
    if sum(len(text) for _, text in units) > MAX_CHARACTERS:
        raise UserError("Extracted text exceeds 2 million characters. Split this document first.")
    if not any(text.strip() for _, text in units):
        raise UserError(
            "No readable text found. "
            + (" ".join(warnings) if warnings else "Add a searchable copy or paste your notes.")
        )
    return units, warnings


def parse_upload(
    name: str, data: bytes, *, isolated: bool = True, use_ocr: bool = True
) -> ParsedDocument:
    name = safe_name(name)
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED:
        raise UserError(
            "Supported files: PDF, DOCX, PPTX, TXT, Markdown, IPYNB, PNG, JPEG and WebP."
        )
    if not data or len(data) > MAX_BYTES:
        raise UserError("Upload a non-empty file no larger than 15 MB.")
    if isolated:
        try:
            proc = subprocess.run(
                [sys.executable, "-m", "modulemate.extract_worker", extension]
                + ([] if use_ocr else ["--no-ocr"]),
                input=data,
                env={
                    k: v
                    for k, v in os.environ.items()
                    if k in {"PATH", "LANG", "LC_ALL", "SYSTEMROOT", "TMPDIR", "TEMP", "TMP"}
                },
                capture_output=True,
                timeout=75
                if use_ocr and extension in {".pdf", ".docx", ".pptx", ".ipynb", *IMAGE_TYPES}
                else 25,
                cwd=Path(__file__).resolve().parent.parent,
                check=False,
            )
        except subprocess.TimeoutExpired as exc:
            raise UserError(
                "This file took too long to read. Try a smaller or text-only version."
            ) from exc
        if proc.returncode != 0:
            raise UserError(
                "The document could not be read within safe resource limits. Try a smaller text-only copy."
            )
        try:
            payload = json.loads(proc.stdout)
        except (ValueError, UnicodeDecodeError) as exc:
            raise UserError("The document reader returned an invalid result.") from exc
        if "error" in payload:
            raise UserError(payload["error"])
        units, warnings = payload["units"], payload["warnings"]
    else:
        units, warnings = _extract(data, extension, use_ocr=use_ocr)
    chunks = []
    seen = set()
    for locator, text in units:
        for heading, part in structured_chunks(text):
            location = locator + (" · " + heading[:160] if heading else "")
            # Keep duplicate text on distinct pages: locators remain meaningful.
            key = (location, part)
            if key not in seen:
                chunks.append(Chunk(location, part))
                seen.add(key)
    if len(chunks) > 3000:
        raise UserError("This file produces more than 3,000 passages. Split it first.")
    text = "\n\n".join(f"[{locator}]\n{text.strip()}" for locator, text in units)
    digest = hashlib.sha256(text.encode()).hexdigest()
    return ParsedDocument(name, digest, text, chunks, warnings)
