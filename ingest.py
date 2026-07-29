"""M1 — Ingestion, extraction and metadata-preserving chunking (FR1).

Pipeline stages:
  0. Legacy conversion .ppt/.doc -> .pptx/.docx  (scripts/convert_corpus.sh)
  1. Extraction     — one text unit per slide (PPTX) or per section (DOCX)
  2. Chunking       — dense units split with overlap; sparse units kept whole
  3. Metadata       — {source_document, slide_or_section, chunk_id} on every chunk

Citation granularity differs by source type, because the two formats expose
different natural anchors:
  * PPTX -> slide number, read directly from slide order.
  * DOCX -> section heading. The assessment brief carries no page numbers once
    extracted, and the .doc -> .docx conversion produces no real Word heading
    styles (every paragraph is 'Plain Text'), so headings are recovered by the
    structural heuristic in `_is_heading`.
"""
from __future__ import annotations

import json
import re
from dataclasses import asdict, dataclass
from pathlib import Path

import docx
from pptx import Presentation

import config

# A separator rule such as "______" or "~~~~~~" carries no meaning.
_SEPARATOR_RE = re.compile(r"^[_~\-=\.\s]+$")


@dataclass
class Chunk:
    """One retrievable passage plus the metadata needed to cite it."""

    chunk_id: str
    source_document: str      # filename stem, as it appears in a citation
    slide_or_section: str     # "slide 12" | "Assessment scheme"
    doc_type: str             # "pptx" | "docx"
    text: str
    n_words: int

    @property
    def citation(self) -> str:
        return f"[Source: {self.source_document}, {self.slide_or_section}]"


# --------------------------------------------------------------------------
# Extraction
# --------------------------------------------------------------------------

def _shape_text(shape) -> list[str]:
    """Recursively pull text out of a shape, including groups and tables."""
    out: list[str] = []
    if shape.shape_type == 6 and hasattr(shape, "shapes"):  # GROUP
        for sub in shape.shapes:
            out.extend(_shape_text(sub))
        return out
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                out.append(line)
        return out
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            line = "".join(run.text for run in para.runs).strip()
            if line:
                out.append(line)
    return out


def _load_ocr(stem: str, ocr_dir: Path) -> dict[str, str]:
    """OCR text for one deck, keyed by 'slide N' (empty dict if not run)."""
    path = ocr_dir / f"{stem}.json"
    if not path.exists():
        return {}
    with open(path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def _strip_boilerplate(per_slide: dict[str, str], min_repeats: int) -> dict[str, str]:
    """Remove header/footer lines that recur across most slides in a deck.

    Every slide in these decks is stamped with a course banner
    ("COMP5511 Artificial Intelligence Concepts") and a page marker
    ("Lecture #8 - 5"). Left in, they dominate the embedding of otherwise
    short slides and make unrelated slides look mutually similar.
    """
    counts: dict[str, int] = {}
    for text in per_slide.values():
        for line in set(text.splitlines()):
            key = line.strip().lower()
            if key:
                counts[key] = counts.get(key, 0) + 1

    banned = {k for k, n in counts.items() if n >= min_repeats}
    cleaned: dict[str, str] = {}
    for locator, text in per_slide.items():
        kept = [
            ln for ln in text.splitlines()
            if ln.strip() and ln.strip().lower() not in banned
        ]
        if kept:
            cleaned[locator] = "\n".join(kept)
    return cleaned


def extract_pptx(path: Path, ocr_dir: Path | None = None) -> list[tuple[str, str]]:
    """Return [(slide_or_section, text)] — one entry per slide.

    Native shape text and OCR text are merged per slide. Native text is
    authoritative; OCR supplements it, and is the only source for the
    image-only slides that make up the majority of these decks.
    """
    prs = Presentation(str(path))
    native: dict[str, str] = {}
    for idx, slide in enumerate(prs.slides, start=1):
        lines: list[str] = []
        for shape in slide.shapes:
            lines.extend(_shape_text(shape))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Speaker notes: {notes}")
        text = "\n".join(dict.fromkeys(lines)).strip()  # de-dup, keep order
        if text:
            native[f"slide {idx}"] = text

    ocr = _load_ocr(path.stem, ocr_dir) if ocr_dir else {}
    if ocr:
        # Boilerplate recurs on nearly every slide; 40% is a safe cut-off that
        # catches banners without removing a genuinely repeated definition.
        ocr = _strip_boilerplate(ocr, max(3, int(0.4 * len(ocr))))

    merged: dict[str, str] = {}
    for locator in sorted(
        set(native) | set(ocr), key=lambda s: int(s.split()[1])
    ):
        parts: list[str] = []
        seen: set[str] = set()
        for source in (native.get(locator, ""), ocr.get(locator, "")):
            for line in source.splitlines():
                key = line.strip().lower()
                if key and key not in seen:
                    seen.add(key)
                    parts.append(line.strip())
        if parts:
            merged[locator] = "\n".join(parts)

    return list(merged.items())


def _is_heading(text: str) -> bool:
    """Structural heading test for a converted .doc with no heading styles.

    A heading is short, is not a separator rule or bullet, does not end in
    sentence punctuation, and — if it contains a colon — has nothing after it
    ("Requirements:" is a heading, "Due: Friday 31 July" is a data line).
    """
    t = text.strip()
    if not t or len(t) > 60:
        return False
    if _SEPARATOR_RE.match(t) or t.startswith(("--", "*", "•")):
        return False
    if t.endswith((".", ";", ",")):
        return False
    if ":" in t and t.partition(":")[2].strip():
        return False
    return True


def extract_docx(path: Path) -> list[tuple[str, str]]:
    """Return [(section_title, text)] — one entry per detected section."""
    doc = docx.Document(str(path))
    sections: list[tuple[str, list[str]]] = []
    current_title = "Introduction"
    current_body: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text or _SEPARATOR_RE.match(text):
            continue
        if _is_heading(text):
            if current_body:
                sections.append((current_title, current_body))
            current_title = text.rstrip(":").strip()
            current_body = []
        else:
            current_body.append(text)

    if current_body:
        sections.append((current_title, current_body))

    # Tables live outside the paragraph flow; attach them to their own section.
    for t_idx, table in enumerate(doc.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                rows.append(line)
        if rows:
            sections.append((f"Table {t_idx}", rows))

    return [(title, "\n".join(body).strip()) for title, body in sections if body]


# --------------------------------------------------------------------------
# Chunking
# --------------------------------------------------------------------------

def chunk_text(text: str, target_words: int, overlap_words: int) -> list[str]:
    """Split on a word budget with overlap. Short units pass through whole."""
    words = text.split()
    if len(words) <= target_words:
        return [text]
    stride = max(1, target_words - overlap_words)
    pieces: list[str] = []
    for start in range(0, len(words), stride):
        window = words[start:start + target_words]
        if not window:
            break
        pieces.append(" ".join(window))
        if start + target_words >= len(words):
            break
    return pieces


def build_chunks(cfg=None) -> list[Chunk]:
    """Run the whole ingestion pipeline over the converted corpus."""
    cfg = cfg or config.load()
    converted = cfg.path("corpus.converted_dir")
    target = cfg.require("index.chunk_target_words")
    overlap = cfg.require("index.chunk_overlap_words")
    min_chars = cfg.require("index.min_chunk_chars")

    ocr_dir = cfg.path("corpus.ocr_dir")

    files = sorted(converted.glob("*.pptx")) + sorted(converted.glob("*.docx"))
    if not files:
        raise FileNotFoundError(
            f"No .pptx/.docx found in {converted}. "
            f"Run ./scripts/convert_corpus.sh first."
        )

    chunks: list[Chunk] = []
    for path in files:
        doc_type = path.suffix.lstrip(".")
        stem = path.stem  # the name used in citations
        units = (
            extract_pptx(path, ocr_dir) if doc_type == "pptx"
            else extract_docx(path)
        )

        for locator, text in units:
            for piece_idx, piece in enumerate(
                chunk_text(text, target, overlap)
            ):
                if len(piece) < min_chars:
                    continue
                slug = re.sub(r"[^a-z0-9]+", "-", locator.lower()).strip("-")
                chunks.append(
                    Chunk(
                        chunk_id=f"{stem}::{slug}::{piece_idx}",
                        source_document=stem,
                        slide_or_section=locator,
                        doc_type=doc_type,
                        text=piece,
                        n_words=len(piece.split()),
                    )
                )
    return chunks


def save_chunks(chunks: list[Chunk], out_path: Path) -> None:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as fh:
        json.dump([asdict(c) for c in chunks], fh, ensure_ascii=False, indent=1)


if __name__ == "__main__":
    cfg = config.load()
    chunks = build_chunks(cfg)
    by_doc: dict[str, int] = {}
    for c in chunks:
        by_doc[c.source_document] = by_doc.get(c.source_document, 0) + 1
    print(f"Extracted {len(chunks)} chunks from {len(by_doc)} documents\n")
    for doc, n in sorted(by_doc.items()):
        print(f"  {n:5d}  {doc}")
    words = [c.n_words for c in chunks]
    print(f"\nwords/chunk: min={min(words)} median={sorted(words)[len(words)//2]} max={max(words)}")
